"""Pablo local adapter. Existing Huashu tools remain unchanged."""
import argparse
import importlib
import json
import os
from pathlib import Path
from xml.sax.saxutils import escape
import shutil
import subprocess
import sys
import tempfile
import unicodedata

ROOT = Path(__file__).resolve().parents[2]
os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(ROOT / '.pablo/browsers')


def run(command):
    result = subprocess.run([str(x) for x in command], cwd=ROOT, text=True,
                            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=180)
    if result.returncode:
        raise RuntimeError(result.stdout[-6000:])
    return result.stdout


def norm(text):
    return ''.join(unicodedata.normalize('NFC', text).split())


def inspect_pptx(file):
    from pptx import Presentation
    deck = Presentation(file)
    slides = []
    def text_of(shape):
        if shape.shape_type == 6:
            return '\n'.join(text_of(s) for s in shape.shapes)
        if shape.has_table:
            return '\n'.join(c.text for r in shape.table.rows for c in r.cells)
        return shape.text if shape.has_text_frame else ''
    for i, slide in enumerate(deck.slides, 1):
        slides.append({'slide': i, 'text': '\n'.join(text_of(s) for s in slide.shapes),
                       'shapes': len(slide.shapes), 'has_notes': slide.has_notes_slide})
    return {'slide_count': len(slides), 'slides': slides,
            'limitation': 'Structure inventory only; does not prove visual quality or preservation of all features.'}


def find_soffice():
    config = ROOT / '.pablo/local.json'
    saved = json.loads(config.read_text()).get('soffice') if config.exists() else None
    candidates = [os.environ.get('PABLO_SOFFICE'), shutil.which('soffice'), saved,
                  '/Applications/LibreOffice.app/Contents/MacOS/soffice']
    return next((p for p in candidates if p and Path(p).is_file() and os.access(p, os.X_OK)), None)


def configure_fonts():
    local = ROOT / '.pablo'
    local.mkdir(exist_ok=True)
    config = local / 'fonts.conf'
    directories = [local / 'fonts', Path('/System/Library/Fonts'), Path('/Library/Fonts'),
                   Path.home() / 'Library/Fonts', Path('/usr/share/fonts')]
    config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                      + ''.join('<dir>' + escape(str(p)) + '</dir>' for p in directories if p.exists())
                      + '<cachedir>' + escape(str(local / 'font-cache')) + '</cachedir></fontconfig>')
    os.environ['FONTCONFIG_FILE'] = str(config)
    os.environ['SAL_FONTPATH'] = str(local / 'fonts')


def render_pptx(file, out):
    import pymupdf
    configure_fonts()
    file, out = Path(file).resolve(), Path(out).resolve()
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError('PPTX renderer unavailable. Install LibreOffice or set PABLO_SOFFICE to its executable.')
    if out.exists():
        raise RuntimeError(f'Use a new output directory: {out}')
    out.mkdir(parents=True)
    try:
        with tempfile.TemporaryDirectory(prefix='pablo-office-') as profile:
            run([soffice, '-env:UserInstallation=' + Path(profile).as_uri(),
                 '--headless', '--convert-to', 'pdf', '--outdir', out, file])
        pdf = out / (file.stem + '.pdf')
        with pymupdf.open(pdf) as document:
            expected = inspect_pptx(file)['slide_count']
            if len(document) != expected:
                raise RuntimeError(f'Rendered page count mismatch: {len(document)} != {expected}')
            for i, page in enumerate(document, 1):
                page.get_pixmap(matrix=pymupdf.Matrix(1.25, 1.25)).save(out / f'{i:02d}.png')
        return pdf
    except Exception:
        (out / 'INCOMPLETE.txt').write_text('Rendering failed. Do not deliver this directory.\n')
        raise


def export_deck(slides, out, expected):
    from playwright.sync_api import sync_playwright
    from pypdf import PdfReader
    slides, out = Path(slides).resolve(), Path(out).resolve()
    files = sorted(slides.glob('*.html'))
    if expected < 1 or len(files) != expected:
        raise RuntimeError(f'Input slide count mismatch: {len(files)} != {expected}')
    if out.exists():
        raise RuntimeError(f'Use a new output directory: {out}')
    out.parent.mkdir(parents=True, exist_ok=True)
    # Publish only after all machine checks pass; partial upstream exports stay temporary.
    with tempfile.TemporaryDirectory(prefix='.pablo-export-', dir=out.parent) as temp:
        stage = Path(temp)
        pptx, pdf = stage / 'deck.pptx', stage / 'deck.pdf'
        log = run(['node', ROOT / 'scripts/export_deck_pptx.mjs', '--slides', slides, '--out', pptx])
        if '✗' in log or '转换失败' in log:
            raise RuntimeError('Upstream converter reported a partial failure:\n' + log)
        inventory = inspect_pptx(pptx)
        if inventory['slide_count'] != expected:
            raise RuntimeError('PPTX slide count mismatch')
        with sync_playwright() as pw:
            browser = pw.chromium.launch()
            try:
                page = browser.new_page(viewport={'width': 1280, 'height': 720})
                errors = []
                page.on('pageerror', lambda e: errors.append(str(e)))
                for i, file in enumerate(files):
                    page.goto(file.as_uri(), wait_until='networkidle')
                    page.evaluate('document.fonts.ready')
                    source = norm(page.locator('body').inner_text())
                    actual = norm(inventory['slides'][i]['text'])
                    if not source or source != actual:
                        raise RuntimeError(f'Text verification failed on {file.name}; inspect text order, generated text, or missing objects.')
                    page.screenshot(path=str(stage / f'html-{i+1:02d}.png'))
                if errors:
                    raise RuntimeError('HTML JavaScript errors: ' + '; '.join(errors))
            finally:
                browser.close()
        run(['node', ROOT / 'scripts/export_deck_pdf.mjs', '--slides', slides,
             '--out', pdf, '--width', '1280', '--height', '720'])
        if len(PdfReader(pdf).pages) != expected:
            raise RuntimeError('PDF page count mismatch')
        (stage / 'verification.json').write_text(json.dumps({
            'expected_slides': expected, 'pptx_slides': inventory['slide_count'],
            'html_pdf_pages': len(PdfReader(pdf).pages), 'text_match': True,
            'visual_review': 'pending; render the PPTX and inspect final images'}, ensure_ascii=False, indent=2))
        shutil.copytree(stage, out)
    print(f'Export checks passed: {out}\nVisual review is still required.')


def doctor():
    failed = False
    for module in ['playwright.sync_api', 'pptx', 'PIL', 'pypdf', 'pymupdf', 'requests']:
        try:
            importlib.import_module(module)
            print('OK Python:', module)
        except ImportError:
            print('MISSING Python:', module)
            failed = True
    try:
        print(run(['node', '-e', "for (const p of ['playwright','sharp','pptxgenjs','pdf-lib']) console.log('OK Node:',p,require.resolve(p));"]).strip())
        print(run(['node', '-e', "const {chromium}=require('playwright'); (async()=>{const b=await chromium.launch();await b.close();console.log('OK Node Chromium');})().catch(e=>{console.error(e.message);process.exit(1)});"]).strip())
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            b.close()
        print('OK Python Chromium')
        if sys.platform == 'darwin':
            print(run(['node', '-e', "const {chromium}=require('playwright'); (async()=>{const b=await chromium.launch({channel:'chrome'});await b.close();console.log('OK Chrome (upstream macOS exporter requirement)');})().catch(e=>{console.error(e.message);process.exit(1)});"]).strip())
    except Exception as e:
        print('MISSING runtime:', e)
        failed = True
    for binary in ['ffmpeg', 'ffprobe', 'soffice']:
        value = find_soffice() if binary == 'soffice' else shutil.which(binary)
        print(('AVAILABLE' if value else 'OPTIONAL/MISSING'), binary, value or '')
    font = ROOT / '.pablo/fonts/NotoSansCJKkr-Regular.otf'
    print('OK Korean font' if font.exists() else 'MISSING Korean font; run ./pablo setup')
    failed = failed or not font.exists()
    print('PPTX visual rendering needs soffice; video export needs ffmpeg/ffprobe.')
    return int(failed)


def smoke():
    from pptx import Presentation
    from pptx.util import Pt
    base = ROOT / 'outputs'
    base.mkdir(exist_ok=True)
    out = Path(tempfile.mkdtemp(prefix='pablo-smoke-', dir=base))
    slides = out / 'slides'
    slides.mkdir()
    for i, title in enumerate(['Pablo 로컬 실행 확인', '편집 가능한 결과물'], 1):
        (slides / f'{i:02d}.html').write_text(f'''<!doctype html><html lang="ko"><meta charset="utf-8"><style>
@font-face{{font-family:'Noto Sans CJK KR';src:url('{(ROOT / '.pablo/fonts/NotoSansCJKkr-Regular.otf').as_uri()}')}}
body{{margin:0;width:960pt;height:540pt;background:#f5f3ee;font-family:'Noto Sans CJK KR',sans-serif;position:relative}}
h1{{position:absolute;left:60pt;top:90pt;font-size:36pt;color:#183e39;margin:0}}
p{{position:absolute;left:60pt;top:180pt;font-size:20pt;color:#293b38;margin:0}}
</style><body><h1>{title}</h1><p>테스트 슬라이드 {i} · 실제 사업 자료가 아닌 환경 검증용입니다.</p></body></html>''')
    export_deck(slides, out / 'export', 2)
    render_pptx(out / 'export/deck.pptx', out / 'rendered')
    deck = Presentation(out / 'export/deck.pptx')
    original = inspect_pptx(out / 'export/deck.pptx')
    for shape in deck.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for text_run in paragraph.runs:
                    text_run.font.size = Pt(30)
            break
    deck.save(out / 'edited.pptx')
    edited = inspect_pptx(out / 'edited.pptx')
    assert [s['text'] for s in original['slides']] == [s['text'] for s in edited['slides']]
    render_pptx(out / 'edited.pptx', out / 'edited-rendered')
    # Real upstream partial failure: one valid and one invalid-sized slide.
    bad = out / 'bad-slides'
    shutil.copytree(slides, bad)
    p = bad / '02.html'
    p.write_text(p.read_text().replace('width:960pt', 'width:100pt'))
    try:
        export_deck(bad, out / 'must-not-exist', 2)
    except RuntimeError as e:
        assert 'partial failure' in str(e), str(e)
        assert not (out / 'must-not-exist').exists()
    else:
        raise RuntimeError('Partial conversion failure was not blocked')
    print(f'SMOKE PASS: export, Korean text, PDF/PPTX counts, PPTX rendering, text-preserving edit, partial failure blocking\n{out}')


def main():
    parser = argparse.ArgumentParser(description='Pablo local setup, conversion and verification; natural-language requests belong in your agent chat.')
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('doctor')
    commands.add_parser('smoke')
    inspect = commands.add_parser('inspect-pptx'); inspect.add_argument('file')
    render = commands.add_parser('render-pptx'); render.add_argument('file'); render.add_argument('--out', required=True)
    export = commands.add_parser('export-deck'); export.add_argument('--slides', required=True); export.add_argument('--out', required=True); export.add_argument('--expected-slides', required=True, type=int)
    args = parser.parse_args()
    if args.command == 'doctor': return doctor()
    if args.command == 'smoke': smoke()
    elif args.command == 'inspect-pptx': print(json.dumps(inspect_pptx(args.file), ensure_ascii=False, indent=2))
    elif args.command == 'render-pptx': print(render_pptx(args.file, args.out))
    elif args.command == 'export-deck': export_deck(args.slides, args.out, args.expected_slides)
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as error:
        print(f'Pablo: {error}', file=sys.stderr)
        sys.exit(1)
